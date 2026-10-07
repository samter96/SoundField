"""UCS 분류표(엑셀) → py/app/data/ucs_codes.json (1회성 도구, 런타임은 JSON 만 읽는다).

    python tools/build_ucs_codes.py "UCS v8.2.1 Full List.xlsx"

UCS 이름 규칙 파일은 분류 ID 로 시작한다 — 'VEHCar_...', 'GLASBrk_...'. 분류 ID 는
분류 약어(CatShort, VEH) + 하위분류 약어(Car, Brk) 다. 파일명에 분류 이름 단어(vehicle,
glass)가 없는 경우가 많아서(실측 2026-10-07: 47만 개 중 24.5만 개) 이 대응표로 검색어를
약어로도 넓힌다 (Database.query 1단계). 행마다 [분류, 하위분류, 분류 ID, 분류 약어].
openpyxl 필요.
"""
import json
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "py" / "app" / "data" / "ucs_codes.json"


def main(xlsx: str) -> None:
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r and r[0] == "Category")
    idx = {h: i for i, h in enumerate(rows[head]) if h}
    out = []
    for r in rows[head + 1:]:
        if not r or not r[idx["CatID"]]:
            continue
        out.append([str(r[idx["Category"]]).strip(), str(r[idx["SubCategory"]]).strip(),
                    str(r[idx["CatID"]]).strip(), str(r[idx["CatShort"]]).strip()])
    OUT.write_text(json.dumps({"source": Path(xlsx).name, "rows": out}, ensure_ascii=False,
                              separators=(",", ":")), encoding="utf-8")
    print(f"OK: {len(out)} 분류 -> {OUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
